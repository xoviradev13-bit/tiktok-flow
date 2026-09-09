# -*- coding: utf-8 -*-
"""
Professional Formatter for Vietnamese Software Development Agreement
Mau_Hop_Dong_Phat_Trien_Phan_Mem_Viet_Nam_Comprehensive_v2.docx
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl
import subprocess

# Colors - Corporate Legal Palette
COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D - Deep Corporate Navy
COLOR_PRIMARY = RGBColor(26, 26, 26)    # #1A1A1A - Off Black for high contrast text
COLOR_MUTED = RGBColor(107, 114, 128)   # #6B7280 - Subtle Slate Gray
COLOR_WHITE = RGBColor(255, 255, 255)   # #FFFFFF

HEX_NAVY = "1B365D"
HEX_LIGHT_BG = "F8FAFC"
HEX_ROW_ALT = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F8FAFC"
HEX_CALLOUT_BORDER = "1B365D"

TOTAL_WIDTH_DXA = 9072  # 160mm (A4 210mm - 30mm left - 20mm right)

def configure_table_pr(table, border_color=HEX_BORDER):
    """Configure table-level properties and subtle borders."""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))

def configure_borderless_table_pr(table):
    """Configure completely borderless table."""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))

def format_cell(cell, width_dxa, fill_hex="FFFFFF", top_m=100, bot_m=100, left_m=120, right_m=120, v_align="center"):
    """Replace tcPr with strict schema-compliant OpenXML properties."""
    tc = cell._tc
    tcPr_xml = (
        f'<w:tcPr {nsdecls("w")}>'
        f'<w:tcW w:w="{width_dxa}" w:type="dxa"/>'
        f'<w:shd w:fill="{fill_hex}"/>'
        f'<w:tcMar>'
        f'<w:top w:w="{top_m}" w:type="dxa"/>'
        f'<w:bottom w:w="{bot_m}" w:type="dxa"/>'
        f'<w:left w:w="{left_m}" w:type="dxa"/>'
        f'<w:right w:w="{right_m}" w:type="dxa"/>'
        f'</w:tcMar>'
        f'<w:vAlign w:val="{v_align}"/>'
        f'</w:tcPr>'
    )
    old_tcPr = tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
    if old_tcPr is not None:
        tc.remove(old_tcPr)
    tc.insert(0, parse_xml(tcPr_xml))

def format_callout_cell(cell, width_dxa=TOTAL_WIDTH_DXA):
    """Format single cell callout box with strict schema-compliant OpenXML properties."""
    tc = cell._tc
    tcPr_xml = (
        f'<w:tcPr {nsdecls("w")}>'
        f'<w:tcW w:w="{width_dxa}" w:type="dxa"/>'
        f'<w:tcBorders>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
        f'<w:shd w:fill="{HEX_CALLOUT_BG}"/>'
        f'<w:tcMar>'
        f'<w:top w:w="160" w:type="dxa"/>'
        f'<w:bottom w:w="160" w:type="dxa"/>'
        f'<w:left w:w="200" w:type="dxa"/>'
        f'<w:right w:w="200" w:type="dxa"/>'
        f'</w:tcMar>'
        f'<w:vAlign w:val="center"/>'
        f'</w:tcPr>'
    )
    old_tcPr = tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
    if old_tcPr is not None:
        tc.remove(old_tcPr)
    tc.insert(0, parse_xml(tcPr_xml))

def make_row_cant_split(row):
    """Prevent table row from splitting across pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    """Mark table row as repeat header on subsequent pages."""
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_field(paragraph, field_name):
    """Insert dynamic Word field (PAGE, NUMPAGES)."""
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{field_name}"/>')
    paragraph._p.append(fldSimple)

def main():
    source_path = "Mau_Hop_Dong_Phat_Trien_Phan_Mem_Viet_Nam_Comprehensive_v2_BACKUP.docx"
    output_path = "Mau_Hop_Dong_Phat_Trien_Phan_Mem_Viet_Nam_Comprehensive_v2.docx"
    
    print("Reading original document elements...")
    src_doc = docx.Document(source_path)
    
    # Extract sequence of elements (Paragraphs & Tables)
    elements = []
    for child in src_doc.element.body:
        if isinstance(child, CT_P):
            p = docx.text.paragraph.Paragraph(child, src_doc)
            elements.append(('P', p))
        elif isinstance(child, CT_Tbl):
            t = docx.table.Table(child, src_doc)
            elements.append(('TBL', t))
            
    print(f"Total elements extracted: {len(elements)}")
    
    # Create new clean target document
    new_doc = docx.Document()
    
    # Configure A4 Page Size & Standard Vietnamese Margins
    section = new_doc.sections[0]
    section.page_width = Inches(8.27)    # 210mm (A4)
    section.page_height = Inches(11.69)  # 297mm (A4)
    section.top_margin = Inches(0.79)    # 2.0 cm
    section.bottom_margin = Inches(0.79) # 2.0 cm
    section.left_margin = Inches(1.18)   # 3.0 cm (Binding margin)
    section.right_margin = Inches(0.79)  # 2.0 cm
    section.different_first_page_header_footer = True
    
    # Configure Normal Style default font
    normal_style = new_doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = COLOR_PRIMARY
    
    # Setup Headers & Footers
    # 1. Subsequent Pages Header
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(4)
    
    pPr = hp._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="4" w:space="4" w:color="{HEX_BORDER}"/></w:pBdr>')
    pPr.append(pBdr)
    tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="{TOTAL_WIDTH_DXA}"/></w:tabs>')
    pPr.append(tabs)
    
    hrun1 = hp.add_run("HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM")
    hrun1.font.name = 'Times New Roman'
    hrun1.font.size = Pt(9)
    hrun1.font.italic = True
    hrun1.font.color.rgb = COLOR_MUTED
    
    hrun_tab = hp.add_run("\tSố: [●]/HĐPTPM/[NĂM]")
    hrun_tab.font.name = 'Times New Roman'
    hrun_tab.font.size = Pt(9)
    hrun_tab.font.italic = True
    hrun_tab.font.color.rgb = COLOR_MUTED
    
    # 2. First Page Footer
    first_footer = section.first_page_footer
    ffp = first_footer.paragraphs[0]
    ffp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ff_run = ffp.add_run("Trang 1 / ")
    ff_run.font.name = 'Times New Roman'
    ff_run.font.size = Pt(9)
    ff_run.font.color.rgb = COLOR_MUTED
    add_field(ffp, "NUMPAGES")
    for r in ffp.runs[1:]:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        r.font.color.rgb = COLOR_MUTED
        
    # 3. Subsequent Pages Footer
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fp.paragraph_format.space_before = Pt(4)
    fp_pPr = fp._p.get_or_add_pPr()
    fp_pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:top w:val="single" w:sz="4" w:space="4" w:color="{HEX_BORDER}"/></w:pBdr>')
    fp_pPr.append(fp_pBdr)
    fp_tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="{TOTAL_WIDTH_DXA}"/></w:tabs>')
    fp_pPr.append(fp_tabs)
    
    frun1 = fp.add_run("Hợp đồng Phát triển Phần mềm | Bảo mật")
    frun1.font.name = 'Times New Roman'
    frun1.font.size = Pt(9)
    frun1.font.italic = True
    frun1.font.color.rgb = COLOR_MUTED
    
    frun2 = fp.add_run("\tTrang ")
    frun2.font.name = 'Times New Roman'
    frun2.font.size = Pt(9)
    frun2.font.color.rgb = COLOR_MUTED
    
    add_field(fp, "PAGE")
    
    frun3 = fp.add_run(" / ")
    frun3.font.name = 'Times New Roman'
    frun3.font.size = Pt(9)
    frun3.font.color.rgb = COLOR_MUTED
    
    add_field(fp, "NUMPAGES")
    for r in fp.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9)
        r.font.color.rgb = COLOR_MUTED

    print("Header and Footer configured successfully.")

    # Process all elements
    elem_idx = 0
    total = len(elements)
    tbl_counter = 0

    while elem_idx < total:
        etype, obj = elements[elem_idx]
        
        if etype == 'P':
            text = obj.text.strip()
            
            # 1. Skip completely empty paragraphs
            if not text:
                elem_idx += 1
                continue
                
            # 2. Quốc hiệu
            if "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in text:
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_PRIMARY
                elem_idx += 1
                continue
                
            # 3. Tiêu ngữ
            if "Độc lập - Tự do - Hạnh phúc" in text:
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run("Độc lập - Tự do - Hạnh phúc")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
                r.font.color.rgb = COLOR_PRIMARY
                elem_idx += 1
                continue
                
            # 4. Underline below Tiêu ngữ (clean centered divider)
            if set(text) <= {'_', ' ', '-'}:
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(16)
                r = p.add_run("──────────────")
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_NAVY
                r.bold = True
                elem_idx += 1
                continue
                
            # 5. Title of Contract
            if "HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM" in text and elem_idx < 10:
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                r = p.add_run("HỢP ĐỒNG PHÁT TRIỂN PHẦN MỀM")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(16)
                r.font.color.rgb = COLOR_NAVY
                elem_idx += 1
                continue
                
            # 6. Contract Number
            if text.startswith("Số:") and elem_idx < 10:
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(14)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.color.rgb = COLOR_MUTED
                elem_idx += 1
                continue
                
            # 7. Legal premises ("Căn cứ...")
            if text.startswith("Căn cứ"):
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.15
                r = p.add_run(text)
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(50, 50, 50)
                elem_idx += 1
                continue
                
            # 8. Date and location ("Hôm nay, ngày [●]...")
            if text.startswith("Hôm nay, ngày"):
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.2
                r = p.add_run(text)
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_PRIMARY
                elem_idx += 1
                continue
                
            # 9. Party Section Headers (BÊN A, BÊN B)
            if text.startswith("BÊN A - ") or text.startswith("BÊN B - "):
                p = new_doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_NAVY
                elem_idx += 1
                continue
                
            # 10. Recital / Agreement intro
            if text.startswith("Bên A và Bên B sau đây gọi riêng là “Bên”"):
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.2
                r = p.add_run(text)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_PRIMARY
                elem_idx += 1
                continue
                
            # 11. Article Headings ("ĐIỀU X.")
            if text.startswith("ĐIỀU ") and "." in text[:10]:
                p = new_doc.add_paragraph()
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(3.5)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_NAVY
                elem_idx += 1
                continue
                
            # 12. Major Section: PHẦN BỔ SUNG
            if text.startswith("PHẦN BỔ SUNG"):
                new_doc.add_page_break()
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(14)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
                r.font.color.rgb = COLOR_NAVY
                elem_idx += 1
                continue
                
            # 13. Major Section: PHỤ LỤC
            if text.startswith("PHỤ LỤC "):
                if any(x in text for x in ["PHỤ LỤC 01", "PHỤ LỤC 04", "PHỤ LỤC 06"]):
                    new_doc.add_page_break()
                    space_top = Pt(10)
                else:
                    space_top = Pt(16)
                    
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = space_top
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(text)
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(13)
                r.font.color.rgb = COLOR_NAVY
                elem_idx += 1
                continue
                
            # 14. Signatures in Main Contract
            if "ĐẠI DIỆN BÊN A" in text and ("(Ký" in elements[min(elem_idx+1, total-1)][1].text if elem_idx+1 < total and elements[min(elem_idx+1, total-1)][0]=='P' else True):
                while elem_idx < total:
                    cur_t, cur_o = elements[elem_idx]
                    if cur_t == 'P' and "PHẦN BỔ SUNG" in cur_o.text:
                        break
                    if cur_t == 'TBL':
                        break
                    elem_idx += 1
                    
                sig_table = new_doc.add_table(rows=1, cols=2)
                sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                sig_table.autofit = False
                configure_borderless_table_pr(sig_table)
                
                half_dxa = TOTAL_WIDTH_DXA // 2
                row = sig_table.rows[0]
                make_row_cant_split(row)
                
                format_cell(row.cells[0], width_dxa=half_dxa, fill_hex="FFFFFF", top_m=100, bot_m=100, left_m=60, right_m=60)
                format_cell(row.cells[1], width_dxa=half_dxa, fill_hex="FFFFFF", top_m=100, bot_m=100, left_m=60, right_m=60)
                
                # Left: Bên A
                p_a1 = row.cells[0].paragraphs[0]
                p_a1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_a1.paragraph_format.space_before = Pt(14)
                p_a1.paragraph_format.space_after = Pt(2)
                p_a1.paragraph_format.keep_with_next = True
                r = p_a1.add_run("ĐẠI DIỆN BÊN A")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_PRIMARY
                
                p_a2 = row.cells[0].add_paragraph()
                p_a2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_a2.paragraph_format.space_before = Pt(0)
                p_a2.paragraph_format.space_after = Pt(65)
                p_a2.paragraph_format.keep_with_next = True
                r = p_a2.add_run("(Ký, ghi rõ họ tên, chức vụ và đóng dấu nếu có)")
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_MUTED
                
                p_a3 = row.cells[0].add_paragraph()
                p_a3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_a3.paragraph_format.space_before = Pt(0)
                p_a3.paragraph_format.space_after = Pt(6)
                r = p_a3.add_run("[Họ và tên / Chức vụ]")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.color.rgb = COLOR_MUTED
                
                # Right: Bên B
                p_b1 = row.cells[1].paragraphs[0]
                p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_b1.paragraph_format.space_before = Pt(14)
                p_b1.paragraph_format.space_after = Pt(2)
                p_b1.paragraph_format.keep_with_next = True
                r = p_b1.add_run("ĐẠI DIỆN BÊN B")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_PRIMARY
                
                p_b2 = row.cells[1].add_paragraph()
                p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_b2.paragraph_format.space_before = Pt(0)
                p_b2.paragraph_format.space_after = Pt(65)
                p_b2.paragraph_format.keep_with_next = True
                r = p_b2.add_run("(Ký, ghi rõ họ tên, chức vụ và đóng dấu nếu có)")
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_MUTED
                
                p_b3 = row.cells[1].add_paragraph()
                p_b3.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_b3.paragraph_format.space_before = Pt(0)
                p_b3.paragraph_format.space_after = Pt(6)
                r = p_b3.add_run("[Họ và tên / Chức vụ]")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                r.font.color.rgb = COLOR_MUTED
                
                continue
                
            # 15. Signatures in Phụ lục 06 (Biên bản nghiệm thu mẫu)
            if "ĐẠI DIỆN BÊN A:" in text and "ĐẠI DIỆN BÊN B:" in text:
                sig_table = new_doc.add_table(rows=1, cols=2)
                sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                sig_table.autofit = False
                configure_borderless_table_pr(sig_table)
                
                half_dxa = TOTAL_WIDTH_DXA // 2
                row = sig_table.rows[0]
                make_row_cant_split(row)
                format_cell(row.cells[0], width_dxa=half_dxa, fill_hex="FFFFFF", top_m=80, bot_m=80, left_m=60, right_m=60)
                format_cell(row.cells[1], width_dxa=half_dxa, fill_hex="FFFFFF", top_m=80, bot_m=80, left_m=60, right_m=60)
                
                p_a1 = row.cells[0].paragraphs[0]
                p_a1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_a1.paragraph_format.space_before = Pt(10)
                p_a1.paragraph_format.space_after = Pt(2)
                r = p_a1.add_run("ĐẠI DIỆN BÊN A")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                
                p_a2 = row.cells[0].add_paragraph()
                p_a2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_a2.paragraph_format.space_before = Pt(0)
                p_a2.paragraph_format.space_after = Pt(55)
                r = p_a2.add_run("(Ký, ghi rõ họ tên)")
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_MUTED
                
                p_b1 = row.cells[1].paragraphs[0]
                p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_b1.paragraph_format.space_before = Pt(10)
                p_b1.paragraph_format.space_after = Pt(2)
                r = p_b1.add_run("ĐẠI DIỆN BÊN B")
                r.bold = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                
                p_b2 = row.cells[1].add_paragraph()
                p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_b2.paragraph_format.space_before = Pt(0)
                p_b2.paragraph_format.space_after = Pt(55)
                r = p_b2.add_run("(Ký, ghi rõ họ tên)")
                r.italic = True
                r.font.name = 'Times New Roman'
                r.font.size = Pt(10)
                r.font.color.rgb = COLOR_MUTED
                
                elem_idx += 1
                continue
                
            # 16. Drafting Notes Callout Box (GHI CHÚ SOẠN THẢO)
            if "GHI CHÚ SOẠN THẢO" in text:
                note_paragraphs = []
                while elem_idx < total:
                    cur_t, cur_o = elements[elem_idx]
                    if cur_t == 'P':
                        cur_text = cur_o.text.strip()
                        if cur_text:
                            note_paragraphs.append(cur_text)
                    elif cur_t == 'TBL':
                        break
                    elem_idx += 1
                    
                note_tbl = new_doc.add_table(rows=1, cols=1)
                note_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                note_tbl.autofit = False
                configure_borderless_table_pr(note_tbl)
                
                cell = note_tbl.rows[0].cells[0]
                format_callout_cell(cell, width_dxa=TOTAL_WIDTH_DXA)
                
                for n_i, n_text in enumerate(note_paragraphs):
                    if n_i == 0:
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_before = Pt(0)
                        cp.paragraph_format.space_after = Pt(4)
                        r = cp.add_run(n_text)
                        r.bold = True
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(11)
                        r.font.color.rgb = COLOR_NAVY
                    else:
                        cp = cell.add_paragraph()
                        cp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.space_after = Pt(4)
                        cp.paragraph_format.line_spacing = 1.15
                        r = cp.add_run(n_text)
                        r.italic = True
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(10)
                        r.font.color.rgb = RGBColor(75, 85, 99)
                        
                continue
                
            # 17. Bullet / Dash items (– text)
            if text.startswith("–") or text.startswith("-"):
                clean_text = text.lstrip("–- ").strip()
                p = new_doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(3.5)
                p.paragraph_format.line_spacing = 1.2
                
                dash_run = p.add_run("– ")
                dash_run.bold = True
                dash_run.font.name = 'Times New Roman'
                dash_run.font.size = Pt(12)
                dash_run.font.color.rgb = COLOR_NAVY
                
                r = p.add_run(clean_text)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.color.rgb = COLOR_PRIMARY
                
                elem_idx += 1
                continue
                
            # 18. Other standard body paragraphs
            p = new_doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.2
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.color.rgb = COLOR_PRIMARY
            elem_idx += 1
            
        elif etype == 'TBL':
            table_obj = obj
            n_rows = len(table_obj.rows)
            n_cols = len(table_obj.columns)
            
            new_table = new_doc.add_table(rows=n_rows, cols=n_cols)
            new_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            new_table.autofit = False
            
            # Precise widths in dxa summing to 9072 dxa (160mm)
            if tbl_counter in (0, 1) and n_cols == 2:
                # Party info (30% label, 70% value)
                col_widths_dxa = [2720, 6352]
            elif tbl_counter == 2 and n_cols == 4:
                # Scope of work (Hạng mục 22%, Mô tả 48%, Thời hạn 15%, Giá trị 15%)
                col_widths_dxa = [1996, 4356, 1360, 1360]
            elif tbl_counter == 3 and n_cols == 5:
                # Payment (Mốc 8%, Điều kiện 32%, Tỷ lệ 14%, Số tiền 23%, Hạn 23%)
                col_widths_dxa = [726, 2903, 1270, 2086, 2087]
            elif tbl_counter == 4 and n_cols == 4:
                # SLA (Mức độ 18%, Mô tả 42%, Thời gian 20%, Khắc phục 20%)
                col_widths_dxa = [1633, 3810, 1814, 1815]
            elif tbl_counter == 5 and n_cols == 3:
                # IP Assets (Loại 26%, Chủ sở hữu 37%, Cách xử lý 37%)
                col_widths_dxa = [2358, 3357, 3357]
            elif tbl_counter == 6 and n_cols == 5:
                # Acceptance checklist (STT 7%, Hạng mục 23%, Tiêu chí 30%, Kết quả 22%, Ghi chú 18%)
                col_widths_dxa = [635, 2086, 2722, 1996, 1633]
            else:
                unit = TOTAL_WIDTH_DXA // n_cols
                col_widths_dxa = [unit] * n_cols
                col_widths_dxa[-1] += (TOTAL_WIDTH_DXA - sum(col_widths_dxa))
                
            configure_table_pr(new_table, border_color=HEX_BORDER)
            
            for r_idx, row in enumerate(table_obj.rows):
                new_row = new_table.rows[r_idx]
                make_row_cant_split(new_row)
                if r_idx == 0:
                    make_row_header(new_row)
                    
                is_header = (r_idx == 0)
                is_party_table = (tbl_counter in (0, 1))
                
                for c_idx, cell in enumerate(row.cells):
                    new_cell = new_row.cells[c_idx]
                    cell_text = cell.text.strip()
                    c_width = col_widths_dxa[c_idx]
                    
                    # Background shading
                    if is_header:
                        fill_color = HEX_NAVY
                    elif is_party_table and c_idx == 0:
                        fill_color = HEX_ROW_ALT
                    elif r_idx % 2 == 1:
                        fill_color = "FFFFFF"
                    else:
                        fill_color = HEX_LIGHT_BG
                        
                    top_pad = 140 if is_header else 100
                    bot_pad = 140 if is_header else 100
                    format_cell(new_cell, width_dxa=c_width, fill_hex=fill_color, top_m=top_pad, bot_m=bot_pad, left_m=120, right_m=120)
                    
                    p = new_cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(0)
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.line_spacing = 1.15
                    
                    # Text alignment
                    if is_header:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif tbl_counter in (0, 1):
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    elif n_cols >= 4 and c_idx == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif tbl_counter == 3 and c_idx in (2,):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif tbl_counter == 6 and c_idx in (0, 3):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif tbl_counter in (2, 4) and c_idx in (2, 3):
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        
                    r = p.add_run(cell_text)
                    r.font.name = 'Times New Roman'
                    
                    if is_header:
                        r.bold = True
                        r.font.size = Pt(11)
                        r.font.color.rgb = COLOR_WHITE
                    elif is_party_table and c_idx == 0:
                        r.bold = True
                        r.font.size = Pt(10.5)
                        r.font.color.rgb = COLOR_NAVY
                    else:
                        r.font.size = Pt(10.5)
                        r.font.color.rgb = COLOR_PRIMARY
                        
            post_p = new_doc.add_paragraph()
            post_p.paragraph_format.space_before = Pt(4)
            post_p.paragraph_format.space_after = Pt(6)
            post_p.paragraph_format.line_spacing = 1.0
            
            tbl_counter += 1
            elem_idx += 1

    print(f"Total tables formatted: {tbl_counter}")
    new_doc.save(output_path)
    print(f"Successfully saved professionally formatted document to: {output_path}")

    # Validate opening with Microsoft Word COM
    print("Validating with Word COM...")
    res = subprocess.run(['powershell', '-ExecutionPolicy', 'Bypass', '-File', 'test_open.ps1', output_path], capture_output=True, text=True)
    print("Word validation result:", res.stdout.strip())

if __name__ == "__main__":
    main()
