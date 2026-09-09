# -*- coding: utf-8 -*-
"""
Generate complete professional Software Development Agreement with real values
from quote and user parameters.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import subprocess

# Colors - Corporate Legal Palette
COLOR_NAVY = RGBColor(27, 54, 93)       # #1B365D - Deep Corporate Navy
COLOR_PRIMARY = RGBColor(26, 26, 26)    # #1A1A1A - Off Black for crisp text
COLOR_MUTED = RGBColor(107, 114, 128)   # #6B7280 - Subtle Slate Gray
COLOR_WHITE = RGBColor(255, 255, 255)   # #FFFFFF

HEX_NAVY = "1B365D"
HEX_LIGHT_BG = "F8FAFC"
HEX_ROW_ALT = "F1F5F9"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "F8FAFC"
HEX_CALLOUT_BORDER = "1B365D"

TOTAL_WIDTH_DXA = 9072  # 160mm (A4 210mm - 30mm left - 20mm right)

def configure_table_pr(table, border_color=HEX_BORDER):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{border_color}"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{border_color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))

def configure_borderless_table_pr(table):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="none"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))

def format_cell(cell, width_dxa, fill_hex="FFFFFF", top_m=100, bot_m=100, left_m=120, right_m=120, v_align="center"):
    tc = cell._tc
    tcPr_xml = (
        f'<w:tcPr {nsdecls("w")}>'
        f'<w:tcW w:w="{width_dxa}" w:type="dxa"/>'
        f'<w:shd w:fill="{fill_hex}"/>'
        f'<w:tcMar>'
        f'<w:top w:w="{top_m}" w:type="dxa"/>'
        f'<w:bottom w:w="{bot_m}" w:type="dxa"/>'
        f'<w:left w:w="{left_m}" w:type="dxa"/>'
        f'<w:right w:w="{right_m}" w:type="dxa"/>'
        f'</w:tcMar>'
        f'<w:vAlign w:val="{v_align}"/>'
        f'</w:tcPr>'
    )
    old_tcPr = tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
    if old_tcPr is not None:
        tc.remove(old_tcPr)
    tc.insert(0, parse_xml(tcPr_xml))

def format_callout_cell(cell, width_dxa=TOTAL_WIDTH_DXA):
    tc = cell._tc
    tcPr_xml = (
        f'<w:tcPr {nsdecls("w")}>'
        f'<w:tcW w:w="{width_dxa}" w:type="dxa"/>'
        f'<w:tcBorders>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_CALLOUT_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
        f'<w:shd w:fill="{HEX_CALLOUT_BG}"/>'
        f'<w:tcMar>'
        f'<w:top w:w="160" w:type="dxa"/>'
        f'<w:bottom w:w="160" w:type="dxa"/>'
        f'<w:left w:w="200" w:type="dxa"/>'
        f'<w:right w:w="200" w:type="dxa"/>'
        f'</w:tcMar>'
        f'<w:vAlign w:val="center"/>'
        f'</w:tcPr>'
    )
    old_tcPr = tc.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}tcPr')
    if old_tcPr is not None:
        tc.remove(old_tcPr)
    tc.insert(0, parse_xml(tcPr_xml))

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

def add_field(paragraph, field_name):
    fldSimple = parse_xml(f'<w:fldSimple {nsdecls("w")} w:instr="{field_name}"/>')
    paragraph._p.append(fldSimple)

def add_p(doc, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=4, line_spacing=1.2, bold=False, italic=False, font_size=12, color=COLOR_PRIMARY, keep_with_next=False, is_bullet=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    
    if is_bullet:
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        r_dash = p.add_run("– ")
        r_dash.bold = True
        r_dash.font.name = 'Times New Roman'
        r_dash.font.size = Pt(font_size)
        r_dash.font.color.rgb = COLOR_NAVY
        
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.name = 'Times New Roman'
    r.font.size = Pt(font_size)
    r.font.color.rgb = color
    return p

def add_article(doc, number_str, title_str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"ĐIỀU {number_str}. {title_str.upper()}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_NAVY
    return p

def add_party_table(doc, party_label, party_data):
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(12)
    p_header.paragraph_format.space_after = Pt(4)
    p_header.paragraph_format.keep_with_next = True
    r = p_header.add_run(party_label)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_NAVY

    table = doc.add_table(rows=len(party_data) + 1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    configure_table_pr(table)
    
    col_widths = [2720, 6352]
    
    r0 = table.rows[0]
    make_row_cant_split(r0)
    make_row_header(r0)
    for c_i, h_text in enumerate(["Thông tin", "Nội dung"]):
        format_cell(r0.cells[c_i], width_dxa=col_widths[c_i], fill_hex=HEX_NAVY, top_m=140, bot_m=140)
        p = r0.cells[c_i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_WHITE
        
    for r_i, (k, v) in enumerate(party_data):
        row = table.rows[r_i + 1]
        make_row_cant_split(row)
        
        format_cell(row.cells[0], width_dxa=col_widths[0], fill_hex=HEX_ROW_ALT)
        p0 = row.cells[0].paragraphs[0]
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(0)
        r = p0.add_run(k)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_NAVY
        
        fill = "FFFFFF" if r_i % 2 == 1 else HEX_LIGHT_BG
        format_cell(row.cells[1], width_dxa=col_widths[1], fill_hex=fill)
        p1 = row.cells[1].paragraphs[0]
        p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.15
        r = p1.add_run(v)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_PRIMARY
        
    post_p = doc.add_paragraph()
    post_p.paragraph_format.space_before = Pt(4)
    post_p.paragraph_format.space_after = Pt(6)

def add_signature_block(doc, rep_a="[Họ và tên Người đại diện]", rep_b="Nguyễn Tiến Đạt", role_a="[Chức vụ]", role_b="Cá nhân / Bên cung cấp dịch vụ"):
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    configure_borderless_table_pr(sig_table)
    
    half_dxa = TOTAL_WIDTH_DXA // 2
    row = sig_table.rows[0]
    make_row_cant_split(row)
    
    format_cell(row.cells[0], width_dxa=half_dxa, fill_hex="FFFFFF", top_m=100, bot_m=100, left_m=60, right_m=60)
    format_cell(row.cells[1], width_dxa=half_dxa, fill_hex="FFFFFF", top_m=100, bot_m=100, left_m=60, right_m=60)
    
    # Left: Bên A
    p_a1 = row.cells[0].paragraphs[0]
    p_a1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a1.paragraph_format.space_before = Pt(14)
    p_a1.paragraph_format.space_after = Pt(2)
    p_a1.paragraph_format.keep_with_next = True
    r = p_a1.add_run("ĐẠI DIỆN BÊN A")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_PRIMARY
    
    p_a2 = row.cells[0].add_paragraph()
    p_a2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a2.paragraph_format.space_before = Pt(0)
    p_a2.paragraph_format.space_after = Pt(65)
    p_a2.paragraph_format.keep_with_next = True
    r = p_a2.add_run("(Ký, ghi rõ họ tên, chức vụ và đóng dấu nếu có)")
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_MUTED
    
    p_a3 = row.cells[0].add_paragraph()
    p_a3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_a3.paragraph_format.space_before = Pt(0)
    p_a3.paragraph_format.space_after = Pt(6)
    r = p_a3.add_run(f"{rep_a}\n{role_a}" if role_a else rep_a)
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_MUTED
    
    # Right: Bên B
    p_b1 = row.cells[1].paragraphs[0]
    p_b1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b1.paragraph_format.space_before = Pt(14)
    p_b1.paragraph_format.space_after = Pt(2)
    p_b1.paragraph_format.keep_with_next = True
    r = p_b1.add_run("ĐẠI DIỆN BÊN B")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    r.font.color.rgb = COLOR_PRIMARY
    
    p_b2 = row.cells[1].add_paragraph()
    p_b2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b2.paragraph_format.space_before = Pt(0)
    p_b2.paragraph_format.space_after = Pt(65)
    p_b2.paragraph_format.keep_with_next = True
    r = p_b2.add_run("(Ký, ghi rõ họ tên, chức vụ và đóng dấu nếu có)")
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_MUTED
    
    p_b3 = row.cells[1].add_paragraph()
    p_b3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_b3.paragraph_format.space_before = Pt(0)
    p_b3.paragraph_format.space_after = Pt(6)
    r = p_b3.add_run(f"{rep_b}\n{role_b}")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_NAVY

print("Helper functions ready.")
